import os
import io
import re
import uuid
import datetime
from typing import Dict, Any, List

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


class DocumentProcessor:
    def __init__(self):
        pass

    def process_file(self, filename: str, content: bytes, department: str = "", tags: str = "") -> Dict[str, Any]:
        doc_id = str(uuid.uuid4())
        ext = os.path.splitext(filename)[1].lower()
        file_size = len(content)

        pages = []
        full_text = ""
        metadata = {
            "id": doc_id,
            "filename": filename,
            "file_size": file_size,
            "file_type": ext,
            "department": department,
            "tags": tags,
            "upload_timestamp": datetime.datetime.utcnow().isoformat(),
            "parser_used": "Unknown"
        }

        if ext == ".pdf":
            pages, full_text = self._extract_pdf(content)
            metadata["parser_used"] = "PyPDF Engine"

        elif ext == ".docx":
            pages, full_text = self._extract_docx(content)
            metadata["parser_used"] = "python-docx Engine"

        elif ext in [".xlsx", ".xls"]:
            pages, full_text = self._extract_excel(content, ext)
            metadata["parser_used"] = "openpyxl Engine"

        elif ext in [".txt", ".md", ".json", ".csv", ".log"]:
            full_text = content.decode("utf-8", errors="ignore")
            pages = [{"page_num": 1, "text": full_text}]
            metadata["parser_used"] = "Plain Text Parser"

        else:
            # Try UTF-8 decode first, then latin-1
            try:
                full_text = content.decode("utf-8", errors="strict")
            except UnicodeDecodeError:
                full_text = content.decode("latin-1", errors="ignore")
            if not full_text.strip():
                full_text = f"[Binary content extracted from {filename}]"
            pages = [{"page_num": 1, "text": full_text}]
            metadata["parser_used"] = "Fallback Text Extractor"

        cleaned_text = self._clean_text(full_text)

        return {
            "document_id": doc_id,
            "filename": filename,
            "metadata": metadata,
            "pages": pages,
            "full_text": cleaned_text,
            "page_count": len(pages),
            "department": department,
            "tags": tags,
            "file_type": ext
        }

    def _extract_pdf(self, content: bytes):
        pages = []
        full_text_parts = []

        if pypdf and content.startswith(b"%PDF-"):
            try:
                reader = pypdf.PdfReader(io.BytesIO(content))
                for idx, page in enumerate(reader.pages):
                    text = page.extract_text() or ""
                    pages.append({"page_num": idx + 1, "text": text})
                    full_text_parts.append(text)
                if pages:
                    return pages, "\n\n".join(full_text_parts)
            except Exception as e:
                print(f"[DocProcessor] PyPDF error: {e}")

        # Fallback: text decoding for text strings or non-binary PDFs
        decoded = content.decode("latin-1", errors="ignore")
        text_matches = re.findall(r'[\x20-\x7E\s]{4,}', decoded)
        fallback_text = "\n".join(text_matches) if text_matches else decoded
        return [{"page_num": 1, "text": fallback_text}], fallback_text

    def _extract_docx(self, content: bytes):
        if DocxDocument is None:
            text = "[DOCX parsing unavailable — install python-docx]"
            return [{"page_num": 1, "text": text}], text

        try:
            doc = DocxDocument(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

            # Also extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n".join(paragraphs)
            # Split into "pages" every 1000 words roughly
            words = full_text.split()
            pages = []
            page_size = 1000
            for i in range(0, max(1, len(words)), page_size):
                page_text = " ".join(words[i:i + page_size])
                pages.append({"page_num": len(pages) + 1, "text": page_text})

            return pages or [{"page_num": 1, "text": full_text}], full_text
        except Exception as e:
            print(f"[DocProcessor] DOCX error: {e}")
            text = f"[DOCX extraction failed: {e}]"
            return [{"page_num": 1, "text": text}], text

    def _extract_excel(self, content: bytes, ext: str):
        if openpyxl is None:
            text = "[Excel parsing unavailable — install openpyxl]"
            return [{"page_num": 1, "text": text}], text

        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            all_text_parts = []
            pages = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows_text = []
                for row in ws.iter_rows(values_only=True):
                    row_values = [str(cell) for cell in row if cell is not None and str(cell).strip()]
                    if row_values:
                        rows_text.append(" | ".join(row_values))

                sheet_text = f"[Sheet: {sheet_name}]\n" + "\n".join(rows_text)
                pages.append({"page_num": len(pages) + 1, "text": sheet_text})
                all_text_parts.append(sheet_text)

            full_text = "\n\n".join(all_text_parts)
            return pages or [{"page_num": 1, "text": "[Empty Excel file]"}], full_text
        except Exception as e:
            print(f"[DocProcessor] Excel error: {e}")
            text = f"[Excel extraction failed: {e}]"
            return [{"page_num": 1, "text": text}], text

    def _clean_text(self, text: str) -> str:
        text = text.replace('\x00', '')
        text = re.sub(r'\r\n', '\n', text)
        text = re.sub(r'[ \t]+', ' ', text)
        text = re.sub(r'\n{3,}', '\n\n', text)
        return text.strip()
